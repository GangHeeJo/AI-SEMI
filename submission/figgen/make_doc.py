import io
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

out = sys.argv[1]
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(1.5)

st = d.styles["Normal"]
st.font.name = "Malgun Gothic"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
st.paragraph_format.space_after = Pt(3)


def para(text, size=10, bold=False, color=None, after=3):
    p = d.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(after)
    return p


def table(rows, widths, header=True, size=9):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(txt); r.font.size = Pt(size)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            if header and i == 0:
                r.bold = True
    return t


para("2026 전국 대학생 AI 반도체 회로 설계 경진대회 · Digital 트랙 2차", 15, True, "0066CC", 1)
para("결과물 제출 안내 (1쪽): ID, 팀명, 결과물 폴더 구성", 11, False, "5B6B76", 6)

table([
    ["항목", "내용"],
    ["ID", "aiasic26911"],
    ["팀명", "REDRED (서울대학교 전기정보공학부 3인)"],
    ["과제명", "회전하는 이벤트 카메라를 하나의 월드 지도로 옮기는 저전력 회로 (좌표 변환 + 월드 메모리)"],
    ["제출물", "① 발표자료 PPT(19쪽)  ② 발표 영상(10분)  ③ 본 설명 1쪽"],
], [3.0, 14.4])

para("", 4)
para("결과물 폴더 구성", 11.5, True, "1B2630", 3)
table([
    ["폴더 / 파일", "내용"],
    ["submission/", "제출 파일: 발표자료 .pptx, 발표 영상, 본 설명서, 그림 생성 스크립트(figgen/)"],
    ["rtl/", "합성 가능한 RTL. 최상위 aer_tx16_coord_transform_v1_rg.v(행 클록 게이팅 메모리) 및 _blink.v(유휴 클록 차단). 좌표 변환 coord_transform_rmcm.v, 쓰기 경로 world_mem_writer.v, 메모리 world_mem_ff_rg.v, 회전각 추정기 bayes_filter.v"],
    ["tb/", "테스트벤치. tb_aer_tx16_coord_transform_v1_uzh_trace.v(실제 UZH 이벤트 8,503개), _correctness.v(무작위 스트레스), tb_coord_transform_rmcm_correctness.v(전수 16,384), tb_world_mem_ff_rg.v, tb_blink.v, tb_bayes_filter_uzh_trace.v"],
    ["syn/", "Genus 합성·Innovus P&R 스크립트(syn/pnr/flow_signoff/), 합성·타이밍·전력 리포트(syn/reports/)"],
    ["scripts/", "소프트웨어 기준 모델과 검증 스크립트: coord_transform_model.py(오라클), rotation_estimate_model.py, delta_*.py, uzh_*.py 등"],
    ["common_traces_*/, results_delta/", "시뮬레이션 입력 트레이스(실제 센서 영상에서 추출)와 결과 그림(월드 지도, 3D 점구름 등)"],
    ["progress.md", "전체 실험·검증 기록(§1~§217): 시도, 기각, 정정까지 시간순으로 기록"],
], [4.6, 12.8])

para("", 4)
para("재현 방법", 11.5, True, "1B2630", 3)
para("① 시뮬레이션(프로젝트 루트에서): iverilog -g2012 -I rtl -y rtl tb/tb_aer_tx16_coord_transform_v1_uzh_trace.v && vvp a.out  → …_UZH_TRACE_PASS 출력", 9, False, None, 2)
para("② 합성·P&R: bash syn/pnr/flow_signoff/run_flow.sh <작업폴더> aer_tx16_coord_transform_v1_rg 5.0 0.6 <rtl 파일들>  (환경변수 CG=1, DEFS=ICG_LIB), 이어서 run_signoff.sh", 9, False, None, 2)
para("③ 소프트웨어 기준 모델: python scripts/coord_transform_model.py  → self-check PASS (좌표 변환 오라클, RTL 비교용 벡터 생성 포함)", 9, False, None, 6)

para("참고", 11.5, True, "1B2630", 3)
for t in ["대용량 원본 데이터(UZH shapes_rotation 약 500 MB, 센서 .dvs/.h5)와 대회 Q&A 자료는 제출 폴더에서 제외했습니다.",
          "프로젝트 전체 이력은 https://github.com/GangHeeJo/AI-SEMI (main)에서 확인할 수 있습니다.",
          "설계 조건: GPDK045 교육용 PDK, 클록 5 ns, 입출력 지연 0.25 ns, uncertainty 0.10 ns (팀 공통 SDC)."]:
    p = d.add_paragraph(style="List Bullet"); r = p.add_run(t); r.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(1)

d.save(out)
print("saved", out)
