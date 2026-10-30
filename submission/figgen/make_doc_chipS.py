import io
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

out = sys.argv[1]
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(1.5)

st = d.styles["Normal"]
st.font.name = "나눔스퀘어 네오 Regular"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "나눔스퀘어 네오 Regular")
st.paragraph_format.space_after = Pt(3)


def setf(r, bold):
    nm = "나눔스퀘어 네오 Bold" if bold else "나눔스퀘어 네오 Regular"
    r.font.name = nm; r.bold = False
    r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), nm)


def para(text, size=10, bold=False, color=None, after=3):
    p = d.add_paragraph()
    r = p.add_run(text); setf(r, bold); r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(after)
    return p


def table(rows, widths, header=True, size=9):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(txt); r.font.size = Pt(size)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            setf(r, bool(header and i == 0))
    return t


para("2026 전국 대학생 AI 반도체 회로 설계 경진대회 · Digital 트랙 2차", 15, True, "0066CC", 1)
para("결과물 제출 안내 (1쪽): ID, 팀명, 결과물 폴더 구성", 11, False, "5B6B76", 6)

table([
    ["항목", "내용"],
    ["ID", "aiasic26911"],
    ["팀명", "REDRED (서울대학교 전기정보공학부 3인)"],
    ["과제명", "회전하는 이벤트 카메라를 하나의 월드 지도로 옮기는 칩 (칩 S: 센서 워드 → 디코더 → 좌표 변환 8레인 → 월드 메모리 + 칩 안 자세 추정기)"],
    ["제출물", "① 발표자료 PPT(20쪽)  ② 발표 영상(10분)  ③ 본 설명 1쪽"],
], [3.0, 14.4])

para("", 4)
para("결과물 폴더 구성 (서버 /home/aiasic26911/ 기준)", 11.5, True, "1B2630", 3)
table([
    ["폴더 / 파일", "내용"],
    ["hyunsu_round2/code/rtl/", "칩 S RTL(SystemVerilog). 최상위 chip_s_full.sv, 디코더 aer_decoder, 변환 8레인 xf_b8, pix_serializer, world_writer, world_mem + 8 kbit 매크로 mem_bank_8k, 자세 추정기 est_top·est_core·est_plane·est_cordic·est_tok_expand"],
    ["hyunsu_round2/code/tb/", "테스트벤치: tb_chip_full(종단), tb_chip_full_gate(게이트 시뮬), tb_est, tb_chip_s, tb_bank"],
    ["hyunsu_round2/code/results/, code/CANDIDATES.md", "설계·검증 기록 job*.md(job23 RTL, job24 추정기, job25 풀칩 P&R, job26 일반화), 후보 판정표 code/CANDIDATES.md, 정수 기준모델 job24_est.py·job23_golden.py"],
    ["hyunsu_round2/job25/", "풀칩 합성·P&R·signoff·LEC·게이트 시뮬 스크립트(flow/), 산출물 full_5.0/pnr/ (chip_s_full3.gds·.def·_pnr3.v·_quantus.spef·_signoff.sdf, signoff_power3_80.rpt)"],
    ["redred-faer/submission/", "제출 파일: 발표자료 .pptx/.pdf, 본 설명서, 발표 영상, 그림·덱 생성 스크립트(figgen/)"],
    ["redred-faer/scripts/, syn/reports/hyunsu_selfcal/", "자기보정(피벗 상수 알고리즘화)을 정답 있는 합성 장면에서 독립 검증한 스크립트·결과·사전 평가 기준"],
    ["redred-faer/syn/reports/, progress.md", "병행 설계(행 클록 게이팅, blink, hold ECO 등)의 P&R 리포트와 시도·기각 전체 기록(§1~§235)"],
], [4.6, 12.8])

para("", 4)
para("재현 방법", 11.5, True, "1B2630", 3)
para("① 게이트 시뮬(SDF): cd ~/hyunsu_round2/job25 && bash flow/gate_sim.sh roll_B MAXIMUM full_5.0/pnr/chip_s_full_pnr3.v full_5.0/pnr/chip_s_full3_signoff.sdf  (pan_B도 동일, 8 ns 주기)", 9, False, None, 2)
para("② 정수 기준모델: cd ~/hyunsu_round2/code/results && python job24_check.py roll  → ECC 대비 0.467° (센서 원본 .dvs 필요)", 9, False, None, 2)
para("③ 합성·P&R·signoff: job25/flow/ 의 run_syn.sh, run_top.sh, run_signoff3_80.sh, lec_full.sh (절차와 결과는 results/job25_fullchip_pnr.md)", 9, False, None, 6)

para("참고", 11.5, True, "1B2630", 3)
for t in ["대용량 원본 데이터(교수 센서 .dvs 약 9 GB, UZH shapes_rotation 약 500 MB)와 대회 Q&A 자료는 제출 폴더에서 제외했습니다.",
          "칩 S 소스는 서버 ~/hyunsu_round2/code(Git 저장소, 마지막 커밋 2026-10-27), 검증·병행 설계 이력은 https://github.com/GangHeeJo/AI-SEMI (main)에서 확인할 수 있습니다.",
          "설계 조건: GPDK045 교육용 PDK, 출하 클록 8.0 ns(125 MHz), signoff = Quantus 추출 + Innovus 타이밍. 자기보정의 칩 RTL 반영은 미완이며 발표자료 11쪽에 현황을 적었습니다."]:
    p = d.add_paragraph(style="List Bullet"); r = p.add_run(t); setf(r, False); r.font.size = Pt(9)
    p.paragraph_format.space_after = Pt(1)

d.save(out)
print("saved", out)
